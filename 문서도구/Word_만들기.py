"""문서 화면(docs_view/index.html)을 Word 파일 하나로 만든다.

웹 화면은 원본 문서(.md)를 `문서화면_만들기.py`로 변환한 결과이므로, 이 도구는
그 화면을 그대로 읽어 같은 내용·같은 순서로 Word에 옮긴다.

- 글(제목·문단·목록·표·코드)은 Word 서식으로 옮긴다.
- 색깔 상자·단계 그림·구조도는 Chrome으로 선명하게 찍어 그림으로 넣는다.
- 5개 메뉴는 "제목 1", 메뉴 안 구역은 "제목 2" 이하로 넣는다.

사용:
    python 문서도구/Word_만들기.py              # 전체
    python 문서도구/Word_만들기.py --menu how   # 메뉴 하나만 미리보기
필요: python-docx(이 PC에만 설치), Node.js, Chrome 또는 Edge. Word가 있으면 목차·쪽 번호를 미리 계산한다.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import lxml.html
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FONT = "맑은 고딕"
MONO = "Consolas"
NAVY = RGBColor(0x10, 0x2A, 0x43)
TEAL = RGBColor(0x08, 0x7F, 0x8C)
MUTED = RGBColor(0x62, 0x7D, 0x98)
GREEN = RGBColor(0x24, 0x7A, 0x52)
AMBER = RGBColor(0xA7, 0x65, 0x00)
CONTENT_WIDTH_CM = 17.0
MENUS = ("intro", "how", "state", "journey", "ops")
MERGED_ARCHITECTURE = "assets/mentoring/시스템_아키텍처_통합구조.html"

# 색깔 상자와 그림: Word 표로 옮기면 모양이 깨지므로 화면을 찍어 넣는다.
VISUAL_CLASSES = ("hero", "workflow", "flow-pair", "card-grid", "tech-grid", "status-grid", "safety-grid")
VISUAL_SELECTOR = ",".join(f".{name}" for name in VISUAL_CLASSES)
SKIP_CLASSES = {"subnav", "pager"}
TABLE_FONT_BY_COLUMNS = {1: 10, 2: 10, 3: 9.5, 4: 9, 5: 8.5}


def classes(el) -> set[str]:
    return set((el.get("class") or "").split())


def text_of(el) -> str:
    return re.sub(r"\s+", " ", el.text_content()).strip()


# ---------- Word 기본 설정 ----------

def set_fonts(rpr_owner, ascii_font: str = FONT, east_font: str = FONT) -> None:
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in list(fonts.attrib):
        if attribute.endswith("Theme"):
            del fonts.attrib[attribute]
    for key, value in (("w:ascii", ascii_font), ("w:hAnsi", ascii_font), ("w:eastAsia", east_font), ("w:cs", ascii_font)):
        fonts.set(qn(key), value)


def shade(element_pr, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    element_pr.append(shading)


def setup_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.0))
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.size = Pt(10)
    normal.font.color.rgb = NAVY
    set_fonts(normal.element)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25
    sizes = {1: 18, 2: 14, 3: 12, 4: 11, 5: 10.5, 6: 10}
    for level, size in sizes.items():
        style = styles[f"Heading {level}"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = TEAL if level == 2 else NAVY
        set_fonts(style.element)
        style.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    for name in ("Title", "List Bullet", "List Bullet 2", "List Bullet 3"):
        set_fonts(styles[name].element)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE", "1")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    return document


def add_field(paragraph, instruction: str, placeholder: str) -> None:
    def fld(kind: str):
        run = paragraph.add_run()
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind)
        run._r.append(char)
        return run

    fld("begin")
    run = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    run._r.append(text)
    fld("separate")
    paragraph.add_run(placeholder)
    fld("end")


def add_cover(document: Document, updated: str, menu_title: str | None) -> None:
    for _ in range(8):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SCHAT 시스템 문서")
    run.font.size, run.font.bold, run.font.color.rgb = Pt(30), True, NAVY
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("SCHAT 병원 실무지침 AI 시스템 · 현재 버전: SCHAT 버전 3")
    run.font.size, run.font.color.rgb = Pt(13), TEAL
    if menu_title:
        sample = document.add_paragraph()
        sample.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sample.add_run(f"미리보기: {menu_title}")
        run.font.size, run.font.color.rgb = Pt(11), AMBER
    for _ in range(10):
        document.add_paragraph()
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date.add_run(f"최종 업데이트: {updated}")
    run.font.size, run.font.bold, run.font.color.rgb = Pt(12), True, NAVY
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("웹 문서 화면(docs_view)과 같은 원본 문서로 만든 Word 변환본")
    run.font.size, run.font.color.rgb = Pt(9), MUTED

    toc_title = document.add_paragraph()
    toc_title.paragraph_format.page_break_before = True
    run = toc_title.add_run("목차")
    run.font.size, run.font.bold, run.font.color.rgb = Pt(18), True, NAVY
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Word에서 열고 F9를 누르면 목차가 표시됩니다.")


# ---------- 글자 단위 ----------

def add_inline(paragraph, el, fmt: dict | None = None, include_self_text: bool = True) -> None:
    fmt = fmt or {}

    def write(text: str, style: dict) -> None:
        if not text:
            return
        text = re.sub(r"[ \t\r\n]+", " ", text)
        if not any(r.text for r in paragraph.runs):
            text = text.lstrip()
        if not text:
            return
        run = paragraph.add_run(text)
        if style.get("bold"):
            run.bold = True
        if style.get("italic"):
            run.italic = True
        if style.get("color") is not None:
            run.font.color.rgb = style["color"]
        if style.get("size"):
            run.font.size = Pt(style["size"])
        if style.get("code"):
            set_fonts(run._r, MONO, FONT)
            run.font.size = Pt(style.get("size") or 9)
            shade(run._r.get_or_add_rPr(), "EEF2F6")

    if include_self_text:
        write(el.text, fmt)
    for child in el:
        tag = child.tag if isinstance(child.tag, str) else ""
        names = classes(child)
        style = dict(fmt)
        if tag in ("strong", "b", "time"):
            style["bold"] = True
        elif tag in ("em", "i"):
            style["italic"] = True
        elif tag == "code":
            style["code"] = True
        elif tag == "small":
            style.update(size=8.5, color=MUTED)
        elif tag == "a":
            style["color"] = TEAL
        if "status-badge" in names:
            badge = paragraph.add_run("● ")
            badge.font.color.rgb = GREEN if "status-done" in names else AMBER
            badge.bold = True
            fmt_tail = dict(fmt, bold=True, color=GREEN if "status-done" in names else AMBER)
            write(child.tail, fmt_tail)
            continue
        previous = paragraph.runs[-1].text if paragraph.runs else ""
        if tag in ("span", "strong") and previous and previous[-1].isascii() and previous[-1].isalnum() and (child.text or "")[:1].isdigit():
            paragraph.add_run(" ")
        elif tag == "span" and previous and not previous.endswith(" "):
            paragraph.add_run(" ")
        if tag == "br":
            paragraph.add_run().add_break(WD_BREAK.LINE)
        elif tag:
            add_inline(paragraph, child, style)
        write(child.tail, fmt)


# ---------- 블록 단위 ----------

class Converter:
    def __init__(self, document: Document, images: dict, image_dir: Path):
        self.document = document
        self.images = images
        self.image_dir = image_dir
        self.counters: dict[str, int] = {}
        self.page = ""
        self.last_heading = ""

    def heading(self, text: str, level: int) -> None:
        level = max(1, min(level, 6))
        self.document.add_heading(text, level=level)
        self.last_heading = text

    def paragraph(self, el, style: dict | None = None, indent_cm: float = 0.0):
        paragraph = self.document.add_paragraph()
        if indent_cm:
            paragraph.paragraph_format.left_indent = Cm(indent_cm)
        add_inline(paragraph, el, style)
        return paragraph

    def picture(self, file: str, css_width: int) -> None:
        width = min(CONTENT_WIDTH_CM, css_width * 2.54 / 96 * 0.95)
        self.document.add_picture(str(self.image_dir / file), width=Cm(width))
        self.document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def next_visual(self) -> None:
        index = self.counters.get(self.page, 0)
        self.counters[self.page] = index + 1
        shots = self.images["pages"].get(self.page, [])
        if index < len(shots):
            self.picture(shots[index]["file"], shots[index]["width"])

    def list_block(self, el, depth: int = 0) -> None:
        ordered = el.tag == "ol"
        start = int(el.get("start") or 1)
        for number, item in enumerate((c for c in el if c.tag == "li"), start=start):
            paragraph = self.document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6 + depth * 0.6)
            paragraph.paragraph_format.first_line_indent = Cm(-0.5)
            paragraph.paragraph_format.space_after = Pt(2)
            marker = paragraph.add_run(f"{number}. " if ordered else "• ")
            marker.font.color.rgb = TEAL
            item_copy = lxml.html.fromstring(lxml.html.tostring(item, encoding="unicode"))
            for nested in list(item_copy):
                if nested.tag in ("ul", "ol"):
                    nested.drop_tree()
            add_inline(paragraph, item_copy)
            for nested in item:
                if nested.tag in ("ul", "ol"):
                    self.list_block(nested, depth + 1)

    def table(self, el) -> None:
        rows = [row for row in el.iter("tr")]
        if not rows:
            return
        grid = [[cell for cell in row if cell.tag in ("th", "td")] for row in rows]
        columns = max(len(row) for row in grid)
        longest = max((len(text_of(cell)) for row in grid for cell in row), default=0)
        size = TABLE_FONT_BY_COLUMNS.get(columns, 8)
        if columns >= 4 and longest > 160:
            size -= 0.5
        table = self.document.add_table(rows=len(grid), cols=columns)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        tbl_pr = table._tbl.tblPr
        width = OxmlElement("w:tblW")
        width.set(qn("w:w"), "5000")
        width.set(qn("w:type"), "pct")
        for old in tbl_pr.findall(qn("w:tblW")):
            tbl_pr.remove(old)
        tbl_pr.append(width)
        header_rows = {index for index, row in enumerate(rows) if row.getparent().tag == "thead" or all(c.tag == "th" for c in grid[index])}
        for r, row in enumerate(grid):
            tr_pr = table.rows[r]._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
            if r in header_rows:
                repeat = OxmlElement("w:tblHeader")
                tr_pr.append(repeat)
            for c in range(columns):
                cell = table.cell(r, c)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                if c < len(row):
                    add_inline(paragraph, row[c], {"bold": r in header_rows})
                for run in paragraph.runs:
                    if run.font.size is None or run.font.size.pt > size:
                        run.font.size = Pt(size)
                if r in header_rows:
                    shade(cell._tc.get_or_add_tcPr(), "E8F5F6")
                if r in header_rows or r <= max(header_rows, default=-1) + 1:
                    # 머리글 줄이 쪽 끝에 홀로 남지 않게 첫 내용 줄과 붙여 둔다.
                    paragraph.paragraph_format.keep_with_next = True
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)

    def pre(self, el) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.3)
        paragraph.paragraph_format.line_spacing = 1.0
        shade(paragraph._p.get_or_add_pPr(), "F1F4F7")
        for index, line in enumerate(el.text_content().rstrip("\n").split("\n")):
            if index:
                paragraph.add_run().add_break(WD_BREAK.LINE)
            run = paragraph.add_run(line)
            set_fonts(run._r, MONO, FONT)
            run.font.size = Pt(8.5)

    def blockquote(self, el) -> None:
        for child in el:
            paragraph = self.paragraph(child, {"color": NAVY}, indent_cm=0.4)
            p_pr = paragraph._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            for key, value in (("w:val", "single"), ("w:sz", "18"), ("w:space", "6"), ("w:color", "087F8C")):
                left.set(qn(key), value)
            border.append(left)
            p_pr.append(border)
            shade(p_pr, "F0F7F8")

    def history(self, el) -> None:
        items = el.findall(".//article")
        table = lxml.html.fromstring("<table><thead><tr><th>날짜</th><th>기록</th><th>요약</th></tr></thead><tbody></tbody></table>")
        body = table.find("tbody")
        for item in items:
            time = item.find(".//time")
            title = item.find(".//strong")
            summary = item.find(".//p")
            row = lxml.html.fromstring("<tr><td></td><td></td><td></td></tr>")
            cells = row.findall("td")
            cells[0].text = text_of(time) if time is not None else ""
            cells[1].text = text_of(title) if title is not None else ""
            cells[2].text = text_of(summary) if summary is not None else ""
            body.append(row)
        self.table(table)

    def diagram(self, el) -> None:
        caption = el.find(".//figcaption")
        if caption is not None:
            strong = caption.find(".//strong")
            span = caption.find(".//span")
            paragraph = self.document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            if strong is not None:
                paragraph.add_run(text_of(strong)).bold = True
            if span is not None:
                note = paragraph.add_run(f"  {text_of(span)}")
                note.font.color.rgb = MUTED
                note.font.size = Pt(9)
        image = el.find(".//img")
        name = Path(image.get("src", "")).name if image is not None else ""
        shot = self.images.get("svgs", {}).get(name)
        if shot:
            self.picture(shot["file"], shot["width"])

    def merged_architecture(self, el, level: int) -> None:
        caption = el.find(".//figcaption")
        if caption is not None:
            strong = caption.find(".//strong")
            self.heading(text_of(strong) if strong is not None else "통합 아키텍처", level + 1)
            span = caption.find(".//span")
            if span is not None:
                note = lxml.html.fromstring(f"<p>{lxml.html.tostring(span, encoding='unicode')}</p>")
                for link in note.findall(".//a"):
                    link.drop_tree()
                text = text_of(note).rstrip(" ·")
                self.paragraph(lxml.html.fromstring(f"<p>{text}</p>"), {"color": MUTED})
        for shot in self.images.get("merged", []):
            self.picture(shot["file"], shot["width"])

    def handover(self, el, base: int) -> None:
        shift = 0
        first = True
        for child in el:
            tag = child.tag if isinstance(child.tag, str) else ""
            match = re.fullmatch(r"h([2-6])", tag)
            if match:
                text = text_of(child)
                n = int(match.group(1))
                if first and n == 2 and text == self.last_heading:
                    shift = -1
                    first = False
                    continue
                self.heading(text, base + (n - 1) + shift)
                first = False
                continue
            first = False
            self.block(child, base)

    def children(self, el, level: int) -> int:
        current = level
        for child in el:
            current = self.block(child, current, container_level=level)
        return current

    def block(self, el, level: int, container_level: int | None = None) -> int:
        """Writes one element. Returns the heading level that applies to the
        following siblings (a heading inside a container changes it)."""
        if not isinstance(el.tag, str):
            return level
        tag, names = el.tag, classes(el)
        parent_level = level if container_level is None else container_level
        if names & SKIP_CLASSES or tag in ("script", "style", "nav", "hr", "iframe"):
            return level
        if el.get("id") == "merged-architecture":
            self.merged_architecture(el, parent_level)
            return level
        if names & set(VISUAL_CLASSES):
            self.next_visual()
            return level
        if tag == "figure" and "mentoring-diagram" in names:
            self.diagram(el)
            return level
        if tag == "h2":
            self.heading(text_of(el), 1)
            return 1
        if tag in ("h3", "h4", "h5"):
            self.heading(text_of(el), parent_level + 1)
            return parent_level + 1
        if tag == "p":
            style = {"color": MUTED} if names & {"lead", "muted", "footer-note"} else None
            if text_of(el):
                self.paragraph(el, style)
            return level
        if tag in ("ul", "ol"):
            if "chips" in names:
                paragraph = self.document.add_paragraph()
                paragraph.add_run(" · ".join(text_of(li) for li in el if li.tag == "li")).bold = True
            else:
                self.list_block(el)
            return level
        if tag == "table":
            self.table(el)
            return level
        if tag == "pre":
            self.pre(el)
            return level
        if tag == "blockquote":
            self.blockquote(el)
            return level
        if tag == "details":
            summary = el.find("summary")
            if "folder" in names and summary is not None:
                paragraph = self.document.add_paragraph()
                add_inline(paragraph, summary, {"bold": True})
                for child in el:
                    if child is not summary:
                        self.block(child, level)
                return level
            own_level = level + 1
            if summary is not None:
                title = (summary.text or "").strip() or text_of(summary)
                small = summary.find("small")
                if "history-panel" not in names:
                    self.heading(title.lstrip("▶ ").strip(), own_level)
            for child in el:
                if child is summary:
                    continue
                if "history-list" in classes(child):
                    self.history(child)
                else:
                    self.block(child, own_level)
            return level
        if tag == "article" and "handover" in names:
            self.handover(el, level)
            return level
        if tag == "article" and "history-item" in names:
            return level
        if tag in ("div", "section", "header", "article", "main", "figure", "aside"):
            self.children(el, level)
            return level
        if text_of(el):
            self.paragraph(el)
        return level

    def page_div(self, page) -> None:
        self.page = page.get("id")
        for child in page:
            if not isinstance(child.tag, str):
                continue
            if "page-head" in classes(child):
                title = child.find(".//h2")
                self.heading(text_of(title), 1)
                for lead in child.findall("p"):
                    self.paragraph(lead, {"color": MUTED})
                continue
            self.block(child, 1)


def take_pictures(html_path: Path, image_dir: Path) -> dict:
    selectors = json.dumps({"visual": VISUAL_SELECTOR, "merged": MERGED_ARCHITECTURE}, ensure_ascii=False)
    subprocess.run(["node", str(ROOT / "문서도구" / "Word_그림_찍기.mjs"), str(html_path), str(image_dir), selectors], check=True)
    return json.loads((image_dir / "manifest.json").read_text(encoding="utf-8"))


def finish_with_word(docx_path: Path, pdf_path: Path | None) -> bool:
    """Word가 있으면 목차·쪽 번호를 미리 계산해 저장한다(없으면 처음 열 때 계산)."""
    if shutil.which("powershell") is None:
        return False
    export = f"$d.ExportAsFixedFormat('{pdf_path}', 17)" if pdf_path else ""
    script = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
try {{
  $d = $w.Documents.Open('{docx_path}', $false, $false)
  foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
  $d.Save()
  {export}
  $d.Close()
}} finally {{ $w.Quit() }}
"""
    result = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    if result.returncode != 0:
        print("[Word] 목차·쪽 번호 미리 계산을 건너뜁니다. 처음 열 때 F9로 갱신하세요.")
        print(result.stderr[-800:])
    return result.returncode == 0


def strip_update_prompt(docx_path: Path) -> None:
    """Word로 목차를 이미 계산했으면, 열 때마다 묻는 '필드 업데이트' 설정을 뺀다."""
    document = Document(str(docx_path))
    settings = document.settings.element
    for node in settings.findall(qn("w:updateFields")):
        settings.remove(node)
    document.save(str(docx_path))


def build(html_path: Path, output: Path, menu: str | None, pdf_path: Path | None = None) -> Path:
    source = html_path.read_text(encoding="utf-8")
    tree = lxml.html.fromstring(source)
    data_node = tree.find(".//script[@id='schat-overview-data']")
    data = json.loads(data_node.text) if data_node is not None else {}
    updated = data.get("generatedOn", "")
    pages = [tree.find(f".//div[@id='{name}']") for name in MENUS]
    pages = [page for page in pages if page is not None and (menu is None or page.get("id") == menu)]
    if not pages:
        raise SystemExit(f"메뉴를 찾지 못했습니다: {menu}")

    with tempfile.TemporaryDirectory(prefix="schat-word-") as temp:
        image_dir = Path(temp)
        images = take_pictures(html_path, image_dir)
        document = setup_document()
        menu_title = text_of(pages[0].find(".//h2")) if menu else None
        add_cover(document, updated, menu_title)
        converter = Converter(document, images, image_dir)
        for page in pages:
            converter.page_div(page)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output))
    if finish_with_word(output.resolve(), pdf_path.resolve() if pdf_path else None):
        strip_update_prompt(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="SCHAT 문서 화면을 Word 파일로 만들기")
    parser.add_argument("--menu", choices=MENUS, help="메뉴 하나만 미리보기로 만들 때")
    parser.add_argument("--input", type=Path, default=ROOT / "docs_view" / "index.html")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--pdf", type=Path, help="Word가 있을 때 확인용 PDF도 저장")
    args = parser.parse_args()
    folder = ROOT / "docs_view" / "Word_변환본"
    output = args.output or folder / (f"SCHAT_시스템_문서_미리보기_{args.menu}.docx" if args.menu else "SCHAT_시스템_문서.docx")
    build(args.input, output, args.menu, args.pdf)
    print(f"Word 파일을 만들었습니다: {output}")


if __name__ == "__main__":
    main()
