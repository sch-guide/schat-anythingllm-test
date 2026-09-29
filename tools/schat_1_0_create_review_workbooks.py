"""SCHAT 1.0 간호사 검수표와 범위 밖 질문 검증표를 만든다.

병원 원문과 실제 답변이 들어갈 수 있어 Git 제외 폴더에만 저장한다.
"""

from pathlib import Path
import importlib.util

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/02_멘토링/산출물_2026-09-30/로컬전용_원자료"


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def style_sheet(ws, widths):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def nurse_review(cases, document_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "21문항 간호사 검수"
    ws.append([
        "번호", "질문", "기대 문서", "기대 페이지", "SCHAT 답변", "표시된 출처",
        "핵심내용 일치", "숫자/시간/용량 정확", "출처 정확", "불필요 내용 없음",
        "최종 판정", "수정 의견"
    ])
    for number, (_case_id, question, pages) in enumerate(cases, 1):
        ws.append([number, question, document_name, ", ".join(map(str, pages)),
                   "테스트 웹 실행 대기", "테스트 웹 실행 대기", "", "", "", "", "", ""])
    validation = DataValidation(type="list", formula1='"적합,부분적합,부적합"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"K2:K{ws.max_row}")
    style_sheet(ws, {"A": 7, "B": 38, "C": 30, "D": 20, "E": 55, "F": 35,
                     "G": 17, "H": 22, "I": 15, "J": 18, "K": 14, "L": 35})
    note = wb.create_sheet("작성 안내")
    note.append(["항목", "안내"])
    note.append(["현재 상태", "로컬 테스트 웹이 실행되지 않아 실제 Gemini 답변은 아직 채우지 않음"])
    note.append(["사람 입력칸", "G~L열은 간호사가 직접 검수하여 입력"])
    style_sheet(note, {"A": 20, "B": 90})
    path = OUT / "SCHAT_간호사_최종검수.xlsx"
    wb.save(path)
    return path


def outside_review(questions):
    wb = Workbook()
    ws = wb.active
    ws.title = "범위밖 질문 검증"
    ws.append(["번호", "질문", "검색된 근거", "검색 점수", "최종 답변",
               "근거 부족 표시", "일반지식/추측 추가", "판정", "원인/메모"])
    for number, question in enumerate(questions, 1):
        ws.append([number, question, "테스트 웹 실행 대기", "", "테스트 웹 실행 대기", "", "", "", ""])
    validation = DataValidation(type="list", formula1='"적합,부분적합,부적합"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"H2:H{ws.max_row}")
    style_sheet(ws, {"A": 7, "B": 48, "C": 42, "D": 18, "E": 55, "F": 18, "G": 21, "H": 14, "I": 35})
    path = OUT / "SCHAT_범위밖질문_검증.xlsx"
    wb.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    a8 = load_module("a8", ROOT / "tools/schat_a8_current_retrieval_eval.py")
    live = load_module("live", ROOT / "tools/schat_1_0_live_review.py")
    paths = [nurse_review(a8.CASES, a8.DOC), outside_review(live.OUT_OF_SCOPE)]
    for path in paths:
        print(path.name)


if __name__ == "__main__":
    main()
