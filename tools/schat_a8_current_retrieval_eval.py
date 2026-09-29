"""현재 SCHAT 본문 벡터 654개를 대상으로 A8 검색 성능을 재측정한다.

운영 코드와 데이터는 변경하지 않는다. Gemini는 질문 21개의 임베딩 생성에만
사용하고, 병원 원문은 외부로 전송하지 않는다.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sqlite3
import unicodedata
import urllib.request
from collections import Counter
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "docs/02_멘토링/산출물_2026-09-30/로컬전용_원자료"
DOC = "2026실무지침서 (26.7).pdf"

# 2026-09-20에 최종 승인된 21문항을 현재 문서명·페이지에 다시 연결했다.
# 한 질문의 답이 여러 페이지에 걸치면 그 페이지들을 모두 정답으로 인정한다.
CASES = [
    ("UAT-S01", "진정간호 목적은?", [159]),
    ("UAT-S02", "진정절차에 대해 알려줘", [160, 161, 162, 163]),
    ("UAT-S03", "진정 전에 뭘 확인해?", [160, 161, 162]),
    ("UAT-S04", "진정할 때 주의할 점은?", [159, 162, 163, 166]),
    ("UAT-S05", "소아 진정은 몇 분마다 확인해?", [162]),
    ("UAT-S06", "성인은 몇 분마다 모니터링해?", [161]),
    ("UAT-S08", "진정 전·중·후 간호를 순서대로 설명해줘", [160, 161, 162, 163]),
    ("UAT-S09", "성인과 소아 진정 모니터링 차이를 비교해줘", [161, 162]),
    ("UAT-S11", "진정은 어떻게 진행하는 거야?", [160, 161, 162, 163]),
    ("UAT-S12", "그중 진정 후에는 어떻게 관찰해?", [161, 162, 163]),
    ("UAT-S13", "환자에게 진정을 시행하기 전 준비부터 회복 후 관찰까지 지침 근거로 정리해줘", [160, 161, 162, 163]),
    ("UAT-S14", "진정 전 산소포화도 확인 내용이 있나요?", [161, 162]),
    ("UAT-T01", "수혈 전에 확인할 준비사항은?", [117, 118, 121, 123, 124]),
    ("UAT-T02", "수혈 절차를 순서대로 알려줘", list(range(117, 128))),
    ("UAT-T03", "수혈 중 환자 상태는 어떻게 관찰해?", [125, 126, 127]),
    ("UAT-T11", "제제별 투여 시간과 속도를 확인해줘", [114]),
    ("UAT-T12", "수혈 시행 전에 해야 할 일을 정리해줘", [117, 118, 121, 123, 124]),
    ("UAT-T13", "수혈하는 동안 주의 깊게 볼 항목은?", [125, 126, 127]),
    ("UAT-T15", "피 넣기 전에 뭐부터 체크해야 돼?", [117, 124]),
    ("UAT-T16", "수혈간호를 전체적으로 핵심 요약해줘", list(range(113, 131))),
    ("UAT-T17", "그중 이상반응 발생 직후 조치만 다시 알려줘", [128]),
]

STOPWORDS = {"알려줘", "알려주세요", "설명해줘", "설명해주세요", "뭐야", "무엇", "무엇인가요", "관련", "질문", "어떻게", "대해", "대해서", "대하여", "대한", "관해", "관하여"}
PARTICLES = ["에서", "으로", "부터", "까지", "에게", "이란", "란", "을", "를", "은", "는", "이", "가", "에", "로", "와", "과", "의", "도", "만"]
ACTION_STEMS = {"평가", "확인", "설명", "사용", "측정", "비교", "준비", "시행", "검사", "관찰", "기록", "투여", "중단", "처치"}
ENDINGS = ["해주세요", "해줘", "해요", "하나요", "합니까", "하는지", "할까요", "하려면", "해야", "하면", "한다", "하다", "해"]
INTENTS = ["준비사항", "주의사항", "종류", "목적", "절차", "방법", "순서", "준비", "주의", "금기"]

def env_file(path: Path):
    values = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line or line.lstrip().startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def nested_records(value):
    if isinstance(value, dict):
        if "id" in value and "metadata" in value and "values" in value:
            yield value
        for item in value.values():
            yield from nested_records(item)
    elif isinstance(value, list):
        for item in value:
            yield from nested_records(item)


def load_corpus():
    db = sqlite3.connect(ROOT / "server/storage/anythingllm.db")
    live_ids = {str(row[0]) for row in db.execute("select vectorId from document_vectors")}
    found = {}
    for path in (ROOT / "server/storage/vector-cache").glob("*.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        for record in nested_records(data):
            rid = str(record.get("id"))
            if rid in live_ids:
                found[rid] = record
    corpus = []
    for rid, record in found.items():
        md = record.get("metadata") or {}
        if md.get("content_type") == "image_description":
            continue
        text = str(md.get("text") or record.get("document") or record.get("text") or "")
        vector = record.get("values") or []
        if not text or not vector:
            continue
        corpus.append({"id": rid, "text": text, "metadata": md, "vector": vector})
    corpus.sort(key=lambda x: (str(x["metadata"].get("title", "")), int(x["metadata"].get("page") or 0), x["id"]))
    return corpus


def embed_questions(questions):
    env = env_file(ROOT / "docker/.env")
    key = env.get("GEMINI_EMBEDDING_API_KEY") or env.get("GEMINI_API_KEY")
    model = env.get("EMBEDDING_MODEL_PREF") or "gemini-embedding-2"
    if not key:
        raise RuntimeError("Gemini embedding API key가 없습니다.")
    endpoint = "https://generativelanguage.googleapis.com/v1beta/openai/embeddings"
    vectors = []
    requests = 0
    for start in range(0, len(questions), 4):
        batch = questions[start : start + 4]
        body = json.dumps({"model": model, "input": batch}).encode("utf-8")
        request = urllib.request.Request(endpoint, data=body, method="POST", headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"})
        with urllib.request.urlopen(request, timeout=90) as response:
            payload = json.loads(response.read().decode("utf-8"))
        rows = sorted(payload.get("data", []), key=lambda x: x.get("index", 0))
        if len(rows) != len(batch):
            raise RuntimeError("질문 임베딩 응답 개수가 맞지 않습니다.")
        vectors.extend(row["embedding"] for row in rows)
        requests += 1
    return vectors, model, requests


def tokenize(text):
    return re.findall(r"[가-힣]+|[a-z0-9]+(?:\.[0-9]+)?", unicodedata.normalize("NFKC", str(text)).lower())


def strip_particle(token):
    for particle in PARTICLES:
        if token.endswith(particle) and len(token) - len(particle) >= 2:
            return token[: -len(particle)]
    return token


def normalize_action(token):
    for ending in ENDINGS:
        if token.endswith(ending) and token[: -len(ending)] in ACTION_STEMS:
            return token[: -len(ending)]
    return token


def query_terms(text):
    result = []
    for token in tokenize(text):
        token = normalize_action(strip_particle(token))
        split = [token]
        for suffix in sorted(INTENTS, key=len, reverse=True):
            if token != suffix and token.endswith(suffix) and len(token) - len(suffix) >= 2:
                split = [strip_particle(token[: -len(suffix)]), suffix]
                break
        for item in split:
            if len(item) >= 2 and item not in STOPWORDS and item not in result:
                result.append(item)
    return result


def cosine(left, right):
    dot = sum(a * b for a, b in zip(left, right))
    ln = math.sqrt(sum(a * a for a in left)); rn = math.sqrt(sum(b * b for b in right))
    return dot / (ln * rn) if ln and rn else 0.0


def bm25_rank(query, corpus, k1=1.5, b=0.75):
    terms = query_terms(query)
    prepared = []
    for row in corpus:
        tokens = tokenize(row["text"])
        prepared.append((row, tokens, Counter(tokens), re.sub(r"\s+", "", unicodedata.normalize("NFKC", row["text"]).lower())))
    avg = sum(len(x[1]) for x in prepared) / max(len(prepared), 1)
    dfs = {}
    for term in terms:
        dfs[term] = sum(1 for _, tokens, counts, compact in prepared if counts.get(term, 0) or (len(term) >= 3 and term in compact))
    ranked = []
    for row, tokens, counts, compact in prepared:
        score = 0.0
        for term in terms:
            freq = counts.get(term, 0)
            if not freq and len(term) >= 3:
                freq = compact.count(term)
            if not freq:
                continue
            df = dfs[term]
            idf = math.log(1 + (len(prepared) - df + 0.5) / (df + 0.5))
            score += idf * (freq * (k1 + 1)) / (freq + k1 * (1 - b + b * len(tokens) / avg))
        section = str(row["metadata"].get("section") or "").lower()
        title = str(row["metadata"].get("title") or "").lower()
        if terms:
            section_cov = sum(1 for t in terms if t in section) / len(terms)
            title_cov = sum(1 for t in terms if t in title) / len(terms)
            score *= 1 + section_cov + title_cov * 0.25
        if score > 0:
            ranked.append((score, row))
    ranked.sort(key=lambda x: (-x[0], str(x[1]["id"])))
    return ranked


def relevant(row, pages):
    md = row["metadata"]
    return str(md.get("title") or md.get("document_name") or "") == DOC and int(md.get("page") or -1) in pages


def metrics(results):
    hits = []; reciprocal = []
    for case in results:
        first = next((r["rank"] for r in case["rankings"] if r["gold_hit"]), None)
        hits.append(1 if first and first <= 10 else 0)
        reciprocal.append(1 / first if first else 0)
    return {"case_count": len(results), "hit_at_10": sum(hits) / len(hits), "mrr": sum(reciprocal) / len(reciprocal)}


def relabel_existing(output):
    """기존 검색 순위는 그대로 두고 현재 정답 문서·쪽 기준으로만 다시 채점한다."""
    gold_by_id = {case_id: pages for case_id, _question, pages in CASES}
    for engine, cases in output["cases"].items():
        for case in cases:
            pages = gold_by_id[case["case_id"]]
            case["gold_document"] = DOC
            case["gold_pages"] = pages
            for row in case["rankings"]:
                try:
                    page = int(row.get("page"))
                except (TypeError, ValueError):
                    page = -1
                row["gold_hit"] = row.get("document") == DOC and page in pages
        output["aggregate_metrics"][engine] = metrics(cases)
    output["gold"] = {
        "case_count": len(CASES),
        "match_rule": "same document title + page",
    }
    return output


def run():
    corpus = load_corpus()
    if len(corpus) != 654:
        raise RuntimeError(f"본문 벡터가 예상 654개가 아니라 {len(corpus)}개입니다.")
    questions = [q for _, q, _ in CASES]
    embeddings, model, request_count = embed_questions(questions)
    engines = {"vector": [], "bm25": [], "hybrid": []}
    for (case_id, question, pages), qvec in zip(CASES, embeddings):
        vector_full = sorted(((cosine(qvec, row["vector"]), row) for row in corpus), key=lambda x: (-x[0], x[1]["id"]))
        bm25_full = bm25_rank(question, corpus)
        vector_pos = {row["id"]: i + 1 for i, (_, row) in enumerate(vector_full)}
        bm25_pos = {row["id"]: i + 1 for i, (_, row) in enumerate(bm25_full)}
        all_ids = set(vector_pos) | set(bm25_pos)
        by_id = {row["id"]: row for row in corpus}
        hybrid_full = sorted(((0.75 / vector_pos.get(rid, 10**9) + 0.25 / bm25_pos.get(rid, 10**9), by_id[rid]) for rid in all_ids), key=lambda x: (-x[0], x[1]["id"]))
        for engine, ranked in [("vector", vector_full), ("bm25", bm25_full), ("hybrid", hybrid_full)]:
            rows = []
            for rank, (score, row) in enumerate(ranked[:10], 1):
                md = row["metadata"]
                rows.append({"rank": rank, "score": score, "id": row["id"], "document": md.get("title") or md.get("document_name") or "", "page": md.get("page"), "section": md.get("section") or "", "text": row["text"], "gold_hit": relevant(row, pages)})
            engines[engine].append({"case_id": case_id, "question": question, "gold_document": DOC, "gold_pages": pages, "rankings": rows})
    output = {"schema_version": 1, "evaluation_date": "2026-09-29", "corpus": {"body_vector_count": len(corpus)}, "gold": {"case_count": len(CASES), "match_rule": "same document title + page"}, "embedding": {"model": model, "request_count": request_count, "question_count": len(questions), "hospital_source_text_sent": False}, "aggregate_metrics": {key: metrics(value) for key, value in engines.items()}, "cases": engines}
    return output


def save_xlsx(output, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "요약"
    ws.append(["검색 방식", "문항 수", "Hit@10", "MRR", "예전 Hit@10", "예전 MRR"])
    old = {"vector": (0.952380952381, 0.43373015873), "bm25": (0.904761904762, 0.582199546485), "hybrid": (0.952380952381, 0.736961451247)}
    labels = {"vector": "의미 검색", "bm25": "단어 검색", "hybrid": "하이브리드"}
    for key in ["vector", "bm25", "hybrid"]:
        m = output["aggregate_metrics"][key]; ws.append([labels[key], m["case_count"], m["hit_at_10"], m["mrr"], old[key][0], old[key][1]])
    ws = wb.create_sheet("정답_문서쪽")
    ws.append(["문항 ID", "질문", "정답 문서", "정답 쪽"])
    for cid,q,pages in CASES: ws.append([cid,q,DOC,", ".join(map(str,pages))])
    ws = wb.create_sheet("상위10개")
    ws.append(["질문", "정답 문서·쪽", "검색 방식", "순위", "찾아온 문서·쪽", "점수", "조각 원문", "정답 여부", "관련 있음 O/X", "메모"])
    for key in ["vector", "bm25", "hybrid"]:
        for case in output["cases"][key]:
            gold = f'{case["gold_document"]} · p.{",".join(map(str,case["gold_pages"]))}'
            for row in case["rankings"]:
                ws.append([case["question"], gold, labels[key], row["rank"], f'{row["document"]} · p.{row["page"]}', row["score"], row["text"], "O" if row["gold_hit"] else "", "", ""])
    for ws in wb.worksheets:
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="2563EB"); c.alignment=Alignment(horizontal="center",wrap_text=True)
        for col in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=min(55,max(12,max((len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,80)+1)),default=12)+2))
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
    for row in wb["요약"].iter_rows(min_row=2,min_col=3,max_col=6):
        for c in row: c.number_format="0.0%" if c.column in (3,5) else "0.000"
    wb.save(path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    parser.add_argument(
        "--relabel-existing",
        action="store_true",
        help="기존 검색 순위를 재사용하고 정답 문서·쪽 판정만 갱신합니다.",
    )
    args = parser.parse_args(); args.output_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.output_dir / "A8_현재운영_21문항_재평가.json"
    if args.relabel_existing:
        result = relabel_existing(json.loads(json_path.read_text(encoding="utf-8")))
    else:
        result = run()
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    save_xlsx(result, args.output_dir / "A8_현재운영_21문항_재평가.xlsx")
    print(json.dumps({"metrics": result["aggregate_metrics"], "embedding_requests": result["embedding"]["request_count"], "body_vectors": result["corpus"]["body_vector_count"]}, ensure_ascii=False))
