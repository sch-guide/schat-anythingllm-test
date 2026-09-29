"""멘토링 WBS Excel에 SCHAT 1.0 후속 작업을 반영한다."""

from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "docs/02_멘토링/산출물_2026-09-30/SCHAT_WBS_2026-09-29_2026-10-16.xlsx"

ROWS = [
    ["3. 품질 정리", "범위 밖 20문항 테스트", "A9·C1", "", "2026-09-29", "2026-10-02", 30, "테스트 웹 대기"],
    ["3. 품질 정리", "프롬프트 YAML 기준본·fallback", "C2", "", "2026-09-29", "2026-10-07", 70, "Node 회귀 검증 대기"],
    ["3. 품질 정리", "최신 문서만 검색하는 버전 관리", "B1·B2", "", "2026-09-29", "2026-10-07", 30, "설계·백업 완료, 테스트 환경 대기"],
    ["3. 품질 정리", "docs_view 구조도 4종 그림 표시", "D5", "", "2026-09-29", "2026-09-30", 100, "완료"],
    ["3. 품질 정리", "RAGAS 도입 계획·비용 정리", "A8·C1", "", "2026-09-29", "2026-09-30", 100, "완료"],
]

wb = load_workbook(PATH)
ws = wb.active
known = {str(row[1].value) for row in ws.iter_rows(min_row=2) if len(row) > 1}
for row in ROWS:
    if row[1] not in known:
        ws.append(row)
        continue
    for current in ws.iter_rows(min_row=2):
        if current[1].value == row[1]:
            for index, value in enumerate(row, 1):
                current[index - 1].value = value
            break
wb.save(PATH)
print(PATH.name)
